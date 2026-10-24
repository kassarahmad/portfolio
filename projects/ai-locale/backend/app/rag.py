"""RAG locale: estrazione testo, suddivisione in blocchi, embedding e ricerca.

Scelta architetturale: i vettori sono salvati in SQLite e confrontati con numpy.
Per un uso personale (fino a qualche decina di migliaia di blocchi) è veloce,
non richiede un servizio in più e finisce nello stesso backup del database.
Qdrant diventa utile solo con volumi molto più grandi: l'interfaccia `search`
è isolata qui proprio per poterlo sostituire.

Se il server di embedding non è attivo, la ricerca usa un punteggio per parole
chiave (BM25 semplificato) e lo dichiara (`embedded=0`).
"""
from __future__ import annotations

import hashlib
import io
import logging
import math
import re
import uuid
import zipfile
from collections import Counter
from pathlib import Path
from typing import Any

import numpy as np

from .db import Database, now
from .llm import LLMError

log = logging.getLogger("ailocale.rag")

TEXT_EXT = {".txt", ".md", ".markdown", ".py", ".js", ".ts", ".tsx", ".jsx", ".java", ".kt", ".sql",
            ".json", ".csv", ".html", ".css", ".c", ".h", ".cpp", ".hpp", ".cs", ".go", ".rs", ".php",
            ".rb", ".yaml", ".yml", ".xml", ".ini", ".toml", ".log", ".sh", ".ps1", ".swift", ".dart"}
IMAGE_DOC_EXT = {".png", ".jpg", ".jpeg", ".webp"}
ALLOWED_EXT = TEXT_EXT | {".pdf", ".docx", ".xlsx", ".pptx"} | IMAGE_DOC_EXT

CHUNK_CHARS = 1000
CHUNK_OVERLAP = 150
EMBED_BATCH = 16


class DocumentError(Exception):
    pass


def sanitize_filename(name: str) -> str:
    name = Path(name.replace("\\", "/")).name
    name = re.sub(r"[^\w.\- ()àèéìòùÀÈÉÌÒÙ]", "_", name).strip(" .")
    return name[:150] or "documento"


def sanitize_collection(name: str) -> str:
    name = re.sub(r"[^\w\- ]", "", (name or "").strip())[:60]
    return name or "generale"


def validate_upload(filename: str, data: bytes, max_bytes: int) -> str:
    ext = Path(filename).suffix.lower()
    if ext not in ALLOWED_EXT:
        raise DocumentError(f"Formato non supportato: {ext or '(nessuna estensione)'}")
    if not data:
        raise DocumentError("File vuoto")
    if len(data) > max_bytes:
        raise DocumentError(f"File troppo grande (max {max_bytes // (1024 * 1024)} MB)")
    if ext == ".pdf" and not data.startswith(b"%PDF"):
        raise DocumentError("Il file non è un PDF valido")
    if ext == ".docx":
        if not data.startswith(b"PK"):
            raise DocumentError("Il file non è un DOCX valido")
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if "word/document.xml" not in z.namelist():
                    raise DocumentError("Il file non è un DOCX valido")
                total = sum(i.file_size for i in z.infolist())
                if total > 300 * 1024 * 1024:
                    raise DocumentError("DOCX sospetto (decompressione eccessiva)")
        except zipfile.BadZipFile as exc:
            raise DocumentError("Il file non è un DOCX valido") from exc
    if ext in IMAGE_DOC_EXT:
        from .vision import sniff_image
        if not sniff_image(data):
            raise DocumentError("Il file non è un'immagine valida")
    if ext in (".xlsx", ".pptx") and not data.startswith(b"PK"):
        raise DocumentError("File Office non valido")
    if ext in TEXT_EXT and b"\x00" in data[:8192]:
        raise DocumentError("Il file contiene dati binari")
    return ext


def extract_pages(path: Path) -> list[tuple[int | None, str]]:
    ext = path.suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        return [(i + 1, page.extract_text() or "") for i, page in enumerate(reader.pages)]
    if ext == ".xlsx":
        import openpyxl
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets:
            rows = [" | ".join("" if v is None else str(v) for v in row) for row in ws.iter_rows(values_only=True)]
            out.append((None, f"Foglio {ws.title}\n\n" + "\n".join(rows[:5000])))
        return out
    if ext == ".pptx":
        from pptx import Presentation
        prs = Presentation(str(path))
        return [(i, "\n".join(sh.text_frame.text for sh in slide.shapes if getattr(sh, "has_text_frame", False)))
                for i, slide in enumerate(prs.slides, 1)]
    if ext in IMAGE_DOC_EXT:
        return [(None, "")]  # testo ottenuto con OCR in add_document
    if ext == ".docx":
        import docx
        d = docx.Document(str(path))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        return [(None, "\n".join(parts))]
    raw = path.read_bytes()
    try:
        text = raw.decode("utf-8")
    except UnicodeDecodeError:
        text = raw.decode("latin-1")
    return [(None, text)]


def chunk_text(pages: list[tuple[int | None, str]], size: int = CHUNK_CHARS,
               overlap: int = CHUNK_OVERLAP) -> list[tuple[int | None, str]]:
    """Divide per paragrafi rispettando una dimensione massima, con sovrapposizione."""
    chunks: list[tuple[int | None, str]] = []
    for page, text in pages:
        text = re.sub(r"[ \t]+", " ", text).strip()
        if not text:
            continue
        paras = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
        current = ""
        for para in paras:
            while len(para) > size:  # paragrafo enorme: taglio duro
                head, para = para[:size], para[size - overlap:]
                if current:
                    chunks.append((page, current))
                    current = ""
                chunks.append((page, head))
            if len(current) + len(para) + 2 <= size:
                current = f"{current}\n\n{para}" if current else para
            else:
                if current:
                    chunks.append((page, current))
                tail = current[-overlap:] if current else ""
                current = f"{tail}\n\n{para}" if tail else para
        if current:
            chunks.append((page, current))
    return chunks


_WORD = re.compile(r"\w{2,}", re.UNICODE)


def tokenize(text: str) -> list[str]:
    return [w.lower() for w in _WORD.findall(text)]


def keyword_scores(query: str, docs: list[str]) -> list[float]:
    """BM25 semplificato."""
    q = set(tokenize(query))
    if not q or not docs:
        return [0.0] * len(docs)
    toks = [tokenize(d) for d in docs]
    avg = sum(len(t) for t in toks) / max(1, len(toks))
    df = Counter(w for t in toks for w in set(t) if w in q)
    n = len(docs)
    scores = []
    for t in toks:
        tf = Counter(w for w in t if w in q)
        s = 0.0
        for w, f in tf.items():
            idf = math.log(1 + (n - df[w] + 0.5) / (df[w] + 0.5))
            s += idf * f * 2.2 / (f + 1.2 * (0.25 + 0.75 * len(t) / max(1, avg)))
        scores.append(s)
    top = max(scores) or 1.0
    return [s / top for s in scores]


class RAG:
    def __init__(self, db: Database, docs_dir: Path, embedder: Any | None, vision: Any | None = None):
        self.db = db
        self.docs_dir = docs_dir
        self.embedder = embedder  # LLMClient verso il server di embedding (o None)
        self.vision = vision      # per l'OCR di PDF scansionati e immagini

    async def _embed(self, texts: list[str]) -> list[list[float]] | None:
        if self.embedder is None:
            return None
        out: list[list[float]] = []
        try:
            for i in range(0, len(texts), EMBED_BATCH):
                out.extend(await self.embedder.embed(texts[i:i + EMBED_BATCH]))
        except LLMError as exc:
            log.warning("embedding non disponibile: %s", exc)
            return None
        return out

    async def add_document(self, filename: str, data: bytes, collection: str, max_bytes: int) -> dict:
        filename = sanitize_filename(filename)
        collection = sanitize_collection(collection)
        validate_upload(filename, data, max_bytes)
        sha = hashlib.sha256(data).hexdigest()
        existing = self.db.one("SELECT * FROM documents WHERE collection=? AND filename=?", (collection, filename))
        if existing and existing["sha256"] == sha:
            return {**existing, "unchanged": True}

        stored = f"{uuid.uuid4().hex[:10]}_{filename}"
        path = self.docs_dir / stored
        path.write_bytes(data)
        try:
            pages = extract_pages(path)
            pages, ocr_used = await self._ocr_if_needed(path, pages)
        except DocumentError:
            path.unlink(missing_ok=True)
            raise
        except Exception as exc:  # noqa: BLE001
            path.unlink(missing_ok=True)
            raise DocumentError(f"Impossibile leggere il documento: {type(exc).__name__}") from exc
        chunks = chunk_text(pages)
        if not chunks:
            path.unlink(missing_ok=True)
            raise DocumentError("Nessun testo estraibile: è un PDF scansionato o un'immagine e serve l'OCR "
                                "(installa Tesseract o attiva un modello di visione)")
        vectors = await self._embed([c[1] for c in chunks])

        if existing:  # aggiornamento: sostituisce la versione precedente
            self.delete_document(existing["id"])
        n_pages = sum(1 for p, _ in pages if p is not None) or None
        cur = self.db.execute(
            "INSERT INTO documents (collection, filename, stored_name, sha256, size, pages, chunks, embedded, created_at) "
            "VALUES (?,?,?,?,?,?,?,?,?)",
            (collection, filename, stored, sha, len(data), n_pages, len(chunks), 1 if vectors else 0, now()))
        doc_id = cur.lastrowid
        for i, (page, text) in enumerate(chunks):
            blob = np.asarray(vectors[i], dtype=np.float32).tobytes() if vectors else None
            self.db.execute("INSERT INTO chunks (document_id, idx, page, text, embedding) VALUES (?,?,?,?,?)",
                            (doc_id, i, page, text, blob))
        log.info("documento %s indicizzato: %d blocchi, embedding=%s", doc_id, len(chunks), bool(vectors))
        return {**self.db.one("SELECT * FROM documents WHERE id=?", (doc_id,)), "replaced": bool(existing)}

    async def _ocr_if_needed(self, path: Path, pages: list[tuple[int | None, str]]):
        ext = path.suffix.lower()
        engine = await self.vision.ocr_available() if self.vision is not None else None
        if ext in IMAGE_DOC_EXT:
            if not engine:
                return pages, False
            return [(None, await self.vision.ocr_image(path.read_bytes()))], True
        if ext != ".pdf" or not engine:
            return pages, False
        from .vision import render_pdf_page
        out, used = [], False
        for page, text in pages:
            if len(text.strip()) < 20 and page is not None and page <= 80:
                import asyncio
                img = await asyncio.to_thread(render_pdf_page, str(path), page - 1)
                text = await self.vision.ocr_image(img)
                used = True
            out.append((page, text))
        if used:
            log.info("OCR (%s) applicato a %s", engine, path.name)
        return out, used

    def delete_document(self, doc_id: int) -> bool:
        doc = self.db.one("SELECT * FROM documents WHERE id=?", (doc_id,))
        if not doc:
            return False
        self.db.execute("DELETE FROM documents WHERE id=?", (doc_id,))
        (self.docs_dir / doc["stored_name"]).unlink(missing_ok=True)
        return True

    def list_documents(self) -> list[dict]:
        return self.db.query("SELECT id, collection, filename, size, pages, chunks, embedded, created_at "
                             "FROM documents ORDER BY created_at DESC")

    def collections(self) -> list[str]:
        return [r["collection"] for r in self.db.query("SELECT DISTINCT collection FROM documents ORDER BY collection")]

    async def reembed_missing(self) -> int:
        """Calcola gli embedding mancanti (es. documenti caricati a server spento)."""
        docs = self.db.query("SELECT id FROM documents WHERE embedded=0")
        done = 0
        for d in docs:
            rows = self.db.query("SELECT id, text FROM chunks WHERE document_id=? ORDER BY idx", (d["id"],))
            vecs = await self._embed([r["text"] for r in rows])
            if not vecs:
                break
            for r, v in zip(rows, vecs):
                self.db.execute("UPDATE chunks SET embedding=? WHERE id=?",
                                (np.asarray(v, dtype=np.float32).tobytes(), r["id"]))
            self.db.execute("UPDATE documents SET embedded=1 WHERE id=?", (d["id"],))
            done += 1
        return done

    async def search(self, query: str, collections: list[str] | None = None, k: int = 5) -> list[dict]:
        sql = ("SELECT c.id, c.page, c.text, c.embedding, d.filename, d.collection FROM chunks c "
               "JOIN documents d ON d.id = c.document_id")
        params: list[Any] = []
        if collections:
            sql += " WHERE d.collection IN (%s)" % ",".join("?" * len(collections))
            params = list(collections)
        rows = self.db.query(sql, params)
        if not rows:
            return []
        kw = keyword_scores(query, [r["text"] for r in rows])
        sem = [0.0] * len(rows)
        mode = "keyword"
        if any(r["embedding"] for r in rows):
            qv = await self._embed([query])
            if qv:
                q = np.asarray(qv[0], dtype=np.float32)
                q /= (np.linalg.norm(q) or 1.0)
                for i, r in enumerate(rows):
                    if r["embedding"]:
                        v = np.frombuffer(r["embedding"], dtype=np.float32)
                        if v.shape == q.shape:
                            sem[i] = float(np.dot(q, v) / (np.linalg.norm(v) or 1.0))
                mode = "hybrid"
        scored = []
        for i, r in enumerate(rows):
            score = 0.75 * sem[i] + 0.25 * kw[i] if mode == "hybrid" else kw[i]
            scored.append((score, r))
        scored.sort(key=lambda x: -x[0])
        out = []
        for score, r in scored[:k]:
            if score <= 0:
                continue
            out.append({"score": round(score, 4), "filename": r["filename"], "collection": r["collection"],
                        "page": r["page"], "text": r["text"], "mode": mode})
        return out
